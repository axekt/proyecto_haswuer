from pptx import Presentation
from pptx.util import Inches, Pt

# Crear la presentación
prs = Presentation()

# --- Diapositiva 1: Portada ---
slide = prs.slides.add_slide(prs.slide_layouts[0])
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "Análisis de Señales y FFT"
subtitle.text = "Resolución del Laboratorio #1\nAutor: chitiño\nIngeniería Civil en Informática - UCT"

# --- Diapositiva 2: El Desafío ---
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "El Desafío de las Señales"
tf = slide.placeholders[1].text_frame
tf.text = "En Arquitectura de Hardware, las señales rara vez son puras; suelen venir mezcladas con interferencias."
p = tf.add_paragraph()
p.text = "Objetivo: Tomar una señal compleja, simular ruido y analizar sus componentes originales."
p = tf.add_paragraph()
p.text = "Para lograrlo, utilizamos la Transformada Rápida de Fourier (FFT)."

# --- Diapositiva 3: ¿Qué es la FFT? ---
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "¿Qué es la FFT?"
tf = slide.placeholders[1].text_frame
tf.text = "La FFT es un algoritmo matemático que actúa como un 'prisma'."
p = tf.add_paragraph()
p.text = "Convierte una señal del dominio del TIEMPO al dominio de la FRECUENCIA."
p = tf.add_paragraph()
p.text = "Nos permite ver exactamente qué frecuencias (en Hz) componen una onda, separando la señal útil del ruido."

# --- Diapositiva 4: El Código - NumPy ---
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Implementación con NumPy"
tf = slide.placeholders[1].text_frame
tf.text = "Se utilizó NumPy para procesar las matrices de datos de forma eficiente."
p = tf.add_paragraph()
p.text = "np.fft.fft(W12): Calcula la transformada de la onda combinada (1000Hz + 50Hz)."
p = tf.add_paragraph()
p.text = "np.fft.fftfreq(...): Crea el eje X para poder leer los resultados en Hertz (Hz) en nuestro gráfico."

# --- Diapositiva 5: El Código - Matplotlib ---
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Visualización con Matplotlib"
tf = slide.placeholders[1].text_frame
tf.text = "Usamos plt.subplots(2, 2) para crear un tablero de 4 gráficas comparativas."
p = tf.add_paragraph()
p.text = "Gráficas 1 y 2: Visualizan la Onda Original (1000 Hz) y la Onda Ruido (50 Hz) por separado."
p = tf.add_paragraph()
p.text = "Gráfica 3: Muestra la suma caótica de ambas en el tiempo."
p = tf.add_paragraph()
p.text = "Gráfica 4 (FFT): Aísla las frecuencias, revelando picos perfectos en 50 Hz y 1000 Hz."

# --- Diapositiva 6: Conclusiones ---
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Conclusiones del Laboratorio"
tf = slide.placeholders[1].text_frame
tf.text = "La matemática de Fourier logró detectar los componentes ocultos en la señal con total precisión."
p = tf.add_paragraph()
p.text = "Entender esto es vital para el diseño de filtros y la limpieza de datos en sistemas reales."
p = tf.add_paragraph()
p.text = "Python (NumPy + Matplotlib) transforma cálculos complejos en análisis visuales fáciles de interpretar."

# Guardar el archivo
prs.save('Presentacion_FFT_chitino.pptx')
print("¡Listo! Tu archivo Presentacion_FFT_chitino.pptx se ha creado en esta misma carpeta.")